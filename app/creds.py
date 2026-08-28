#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""读写 cred.xlsx 里的 BytePlus 账号凭据。

表结构(第一行为表头，列顺序无关、大小写不敏感):
    EMAIL | UID | AK | SK | STATUS(可选)

STATUS 是软删除标记，**可选**: 没有这一列、或某行该格为空，都算「启用」。
只有明确写了停用值(停用/disabled/0/off/否/no/false)才算停用。这个方向是
故意的 —— 手工加的行不会因为忘了填 STATUS 而凭空消失。

安全约定
--------
本模块读取的是**明文**凭据表。SK 是最高敏感级别的数据，因此:
  * Account.sk 标记了 repr=False —— SK 不会出现在 repr()/f-string/traceback 里；
  * 调用方(server.py)必须保证 SK 永不进入任何 HTTP 响应、日志或异常消息；
  * 对外展示只用 Account.ak_masked（AK 是标识符而非密钥，脱敏后可显示）。

写入约定(add_account / set_enabled)
-----------------------------------
这个文件是整个系统唯一的凭据来源，写坏了等于所有账号一起丢，所以每次写入都是:
  取锁 -> 备份并校验备份可读 -> 原地整体覆盖 -> 重新解析验证 -> 失败则回滚。
细节见 _write_workbook()。
"""

from __future__ import annotations

import datetime
import glob
import io
import os
import shutil
import threading
from dataclasses import dataclass, field
from typing import List, Optional, Tuple

# 项目根目录 = 本文件所在 app/ 的上一级
PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_CRED_PATH = os.path.join(PROJECT_ROOT, "cred.xlsx")

_REQUIRED_COLUMNS = ("EMAIL", "UID", "AK", "SK")
STATUS_COLUMN = "STATUS"

# 写进表格的值。读的时候两边都认，写的时候统一用中文，方便直接在 Excel 里看。
STATUS_ENABLED = "启用"
STATUS_DISABLED = "停用"
_DISABLED_WORDS = frozenset({
    "停用", "禁用", "关闭", "否", "无效",
    "disabled", "disable", "off", "no", "false", "0", "n",
})

# 同一进程内串行化所有写入。多个浏览器标签页同时点「停用」时，两个线程各自
# 读一遍再各自写回，后写的会把先写的覆盖掉(经典 read-modify-write 竞争)。
_WRITE_LOCK = threading.Lock()

# cache/ 里最多留这么多份备份，多了自动删最旧的
MAX_BACKUPS = 10


class CredError(Exception):
    """凭据表读取/校验失败。"""


@dataclass(frozen=True)
class Account:
    email: str
    uid: str
    ak: str
    # repr=False: 防止 SK 通过 repr()/日志/traceback 意外泄漏
    sk: str = field(repr=False)
    row: int = 0
    enabled: bool = True

    @property
    def ak_masked(self) -> str:
        """AK 脱敏显示(AK 是访问标识符，不是密钥)。"""
        if len(self.ak) <= 12:
            return self.ak[:3] + "***"
        return "{}...{}".format(self.ak[:6], self.ak[-4:])

    @property
    def label(self) -> str:
        return "{} ({})".format(self.email, self.uid) if self.email else str(self.uid)

    def __str__(self) -> str:  # 绝不包含 SK
        return self.label


def parse_status(value) -> bool:
    """STATUS 单元格 -> 是否启用。

    **空值 = 启用**。方向是故意选的: 手工往表里加行的人不会记得填 STATUS，
    如果空值当停用，新加的账号会莫名其妙不出现，而且完全看不出原因。
    """
    s = _cell(value).strip().lower()
    return s not in _DISABLED_WORDS if s else True


def _cell(value) -> str:
    """把单元格转成规整的字符串。"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        # Excel 常把 UID 这类整数读成 3001315439.0
        return str(int(value))
    return str(value).strip()


def load_accounts(path: Optional[str] = None,
                  sheet: Optional[str] = None) -> List[Account]:
    """读取凭据表，返回 Account 列表。失败抛 CredError。"""
    path = path or DEFAULT_CRED_PATH
    if not os.path.exists(path):
        raise CredError("凭据文件不存在: {}".format(path))

    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise CredError("缺少依赖 openpyxl，请运行: pip install openpyxl") from exc

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise CredError("无法打开 {}: {}".format(os.path.basename(path), exc)) from exc

    try:
        if sheet is not None:
            if sheet not in wb.sheetnames:
                raise CredError("找不到工作表 '{}'，实际有: {}".format(sheet, wb.sheetnames))
            ws = wb[sheet]
        else:
            ws = wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()

    # ---- 定位表头(跳过顶部空行) ----
    header_idx = None
    for i, row in enumerate(rows):
        if any(_cell(c) for c in row):
            header_idx = i
            break
    if header_idx is None:
        raise CredError("凭据表 '{}' 是空的".format(ws.title))

    header = [_cell(c).upper() for c in rows[header_idx]]
    col = {}
    for name in _REQUIRED_COLUMNS:
        if name not in header:
            raise CredError(
                "凭据表缺少列 '{}'。实际表头: {}".format(name, header or "(空)"))
        col[name] = header.index(name)
    # STATUS 是可选列: 老表没有这一列时全部按启用处理
    status_i = header.index(STATUS_COLUMN) if STATUS_COLUMN in header else None

    # ---- 逐行解析 ----
    accounts: List[Account] = []
    for offset, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        if not any(_cell(c) for c in row):
            continue  # 跳过空行

        def get(name: str) -> str:
            i = col[name]
            return _cell(row[i]) if i < len(row) else ""

        email, uid, ak, sk = get("EMAIL"), get("UID"), get("AK"), get("SK")
        missing = [n for n, v in (("UID", uid), ("AK", ak), ("SK", sk)) if not v]
        if missing:
            raise CredError(
                "cred.xlsx 第 {} 行缺少 {} —— 请补全或删除该行".format(
                    offset, "/".join(missing)))

        status_raw = ""
        if status_i is not None and status_i < len(row):
            status_raw = _cell(row[status_i])

        accounts.append(Account(email=email, uid=uid, ak=ak, sk=sk, row=offset,
                                enabled=parse_status(status_raw)))

    if not accounts:
        raise CredError("凭据表里没有有效账号(只有表头?)")

    # UID 重复会让前端行无法区分，提前拦住
    seen = {}
    for a in accounts:
        if a.uid in seen:
            raise CredError(
                "UID {} 在第 {} 行和第 {} 行重复了".format(a.uid, seen[a.uid], a.row))
        seen[a.uid] = a.row

    return accounts


def find_account(accounts: List[Account], selector: str) -> Optional[Account]:
    """按邮箱 / UID / 序号(1 开始) 查找账号。"""
    s = str(selector).strip().lower()
    for i, a in enumerate(accounts, start=1):
        if s in (a.email.lower(), a.uid.lower(), str(i)):
            return a
    return None


# ===========================================================================
# 写入
# ===========================================================================
# 只有 add_account() 和 set_enabled() 两个入口，两者都必须经过 _write_workbook()。
# 别在别处直接 wb.save(cred.xlsx) —— 备份、回读校验、回滚全在那一个函数里。

def _timestamp() -> str:
    # 带毫秒。只到秒的话同一秒内的两次写入会指向同一个备份文件名，
    # 后一次直接把前一次的备份覆盖掉 —— 连点两下「停用」就少一个回滚点。
    return datetime.datetime.now().strftime("%Y%m%d-%H%M%S-%f")[:-3]


def _prune_backups(backup_dir: str) -> None:
    """只留最近 MAX_BACKUPS 份，多了删最旧的(cache/ 不该无限长大)。"""
    files = sorted(glob.glob(os.path.join(backup_dir, "cred-*.xlsx")))
    for old in files[:-MAX_BACKUPS]:
        try:
            os.remove(old)
        except OSError:
            pass        # 删不掉不是什么大事，别让它挡住写入


def _backup(path: str, backup_dir: str) -> str:
    """备份凭据表，并**验证备份能被解析**。返回备份路径。

    验证这一步不是多余的: 一个存不回去的备份等于没有备份，而我们正要拿它
    当回滚点。备份放 cache/ 而不是凭据表旁边 —— 线上 systemd 只给了 cache/
    和 cred.xlsx 本身写权限，项目目录是只读的(见 deploy/byteplus-billing.service)。
    """
    if not backup_dir:
        raise CredError("没有指定备份目录，拒绝写入凭据表")
    os.makedirs(backup_dir, exist_ok=True)
    dst = os.path.join(backup_dir, "cred-{}.xlsx".format(_timestamp()))
    shutil.copyfile(path, dst)
    try:
        with open(dst, "rb") as fh:
            os.fsync(fh.fileno())      # 确保备份真的落盘，而不是还在页缓存里
    except OSError:
        pass                            # 某些文件系统不支持对只读句柄 fsync
    load_accounts(dst)                  # 解析不了就抛 CredError，写入不会发生
    _prune_backups(backup_dir)
    return dst


def _write_workbook(wb, path: str, backup_dir: str) -> None:
    """把工作簿写回 path。**所有写入的唯一出口。**

    顺序: 备份(并验证) -> 内存里拼出完整字节 -> 一次写入 -> 回读解析 ->
    失败就从备份回滚。

    这里是**原地覆盖**而不是"临时文件 + os.replace"。原子替换需要目录可写，
    而线上 ProtectSystem=strict 只放开了 cred.xlsx 这一个文件，项目目录是
    只读的。所以退而求其次: 内容先在内存里拼完整，落盘只有一次 write，
    并且前面那份已验证的备份就是兜底 —— 真写坏了照着报错里的路径拷回来即可。
    """
    backup = _backup(path, backup_dir)

    buf = io.BytesIO()
    wb.save(buf)                        # 先在内存里生成，避免半截文件
    data = buf.getvalue()

    try:
        with open(path, "wb") as fh:
            fh.write(data)
            fh.flush()
            os.fsync(fh.fileno())
    except OSError as exc:
        raise CredError(
            "写入 {} 失败: {}。原文件可能已损坏，备份在 {}".format(
                os.path.basename(path), exc, backup)) from exc

    # 回读验证: 写完必须还能被 load_accounts() 解析，否则整个系统起不来
    try:
        load_accounts(path)
    except Exception as exc:
        try:
            shutil.copyfile(backup, path)
            rolled = "已从备份回滚"
        except OSError as restore_exc:
            rolled = "回滚也失败了({})，请手动拷回 {}".format(restore_exc, backup)
        raise CredError("写入后校验失败: {}。{}".format(exc, rolled)) from exc


def _header_map(ws) -> Tuple[int, dict]:
    """返回 (表头行号, {列名大写: 列号})，都是 1 起始(openpyxl 的坐标)。"""
    for r in range(1, ws.max_row + 1):
        values = [_cell(ws.cell(row=r, column=c).value)
                  for c in range(1, ws.max_column + 1)]
        if any(values):
            return r, {v.upper(): i + 1 for i, v in enumerate(values) if v}
    raise CredError("凭据表是空的")


def _ensure_status_column(ws, header_row: int, cols: dict) -> int:
    """返回 STATUS 列的列号；老表没有这一列就在末尾补上表头。"""
    if STATUS_COLUMN in cols:
        return cols[STATUS_COLUMN]
    idx = ws.max_column + 1
    ws.cell(row=header_row, column=idx, value=STATUS_COLUMN)
    return idx


def _clean_sheet(ws, header_row: int, cols: dict) -> int:
    """写入前清理工作表，返回**最后一行真实数据**的行号。

    为什么必须清: Excel 会把邮箱自动变成 mailto: 超链接，而 openpyxl 在可写
    模式下(load_workbook 不带 read_only)会把「有超链接、但没有值」的单元格
    **用超链接目标当值创建出来**。表格里删过行之后就会留下这种孤立超链接，
    于是一保存就凭空多出一行 'mailto:xxx@yyy.com' —— 它有内容但没有
    UID/AK/SK，load_accounts 直接报错，写入被回滚，此后再也加不了账号。
    (实测踩到过: 删掉一行后新增/停用全部失败。)

    所以这里做两件事:
      1. 清掉全表超链接。凭据表里的 mailto 链接毫无用处，留着只是隐患。
      2. 认 UID/AK/SK 找出最后一行真实数据(**不认 EMAIL** —— 假行恰好只有
         EMAIL 列有值)，然后把它后面的内容全部清空。
    只清尾部、不动中间行: 中间若真有半填的行，load_accounts 早就报错了，
    根本走不到写入这一步。
    """
    for row in ws.iter_rows():
        for cell in row:
            if cell.hyperlink is not None:
                cell.hyperlink = None
    try:
        ws._hyperlinks = []
    except Exception:      # openpyxl 内部结构变了也不该因此写不进去
        pass

    last = header_row
    for r in range(header_row + 1, ws.max_row + 1):
        if any(_cell(ws.cell(row=r, column=cols[n]).value)
               for n in ("UID", "AK", "SK")):
            last = r

    for r in range(last + 1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            ws.cell(row=r, column=c).value = None
    return last


def _load_for_write(path: str):
    """以可写方式打开工作簿(read_only 模式改不了单元格)。"""
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise CredError("缺少依赖 openpyxl，请运行: pip install openpyxl") from exc
    try:
        return openpyxl.load_workbook(path)
    except Exception as exc:
        raise CredError("无法打开 {}: {}".format(os.path.basename(path), exc)) from exc


def add_account(email: str, uid: str, ak: str, sk: str,
                path: Optional[str] = None,
                backup_dir: Optional[str] = None) -> Account:
    """往凭据表追加一个账号。返回新建的 Account。

    EMAIL 可以为空(load_accounts 只要求 UID/AK/SK)，但 UID/AK/SK 必填。
    UID 和 AK 都要查重: UID 重复会让 load_accounts 直接报错、整个系统起不来；
    AK 重复几乎一定是复制粘贴错了(一把 AK 只属于一个账号)。
    """
    path = path or DEFAULT_CRED_PATH
    email = _cell(email).strip()
    uid, ak, sk = (_cell(v).strip() for v in (uid, ak, sk))

    missing = [n for n, v in (("UID", uid), ("AK", ak), ("SK", sk)) if not v]
    if missing:
        raise CredError("{} 不能为空".format(" / ".join(missing)))

    with _WRITE_LOCK:
        for a in load_accounts(path):
            if a.uid == uid:
                raise CredError("UID {} 已经在第 {} 行了".format(uid, a.row))
            if a.ak == ak:
                raise CredError("这个 AK 已经被 {} 用了(第 {} 行)".format(
                    a.label, a.row))

        wb = _load_for_write(path)
        ws = wb.worksheets[0]
        header_row, cols = _header_map(ws)
        status_col = _ensure_status_column(ws, header_row, cols)
        # 清理并拿到最后一行真实数据，接在它后面
        new_row = _clean_sheet(ws, header_row, cols) + 1

        ws.cell(row=new_row, column=cols["EMAIL"], value=email)
        # UID 一律写成字符串: 写成数字的话 Excel 可能显示成科学计数法，
        # 而它是标识符不是数值，不需要参与任何计算
        ws.cell(row=new_row, column=cols["UID"], value=uid)
        ws.cell(row=new_row, column=cols["AK"], value=ak)
        ws.cell(row=new_row, column=cols["SK"], value=sk)
        ws.cell(row=new_row, column=status_col, value=STATUS_ENABLED)

        _write_workbook(wb, path, backup_dir)

    for a in load_accounts(path):
        if a.uid == uid:
            return a
    raise CredError("写入成功但没读回新账号，请检查 {}".format(path))


def set_enabled(uid: str, enabled: bool,
                path: Optional[str] = None,
                backup_dir: Optional[str] = None) -> Account:
    """启用/停用某个账号(软删除)。状态没变化时不写文件。"""
    path = path or DEFAULT_CRED_PATH
    uid = _cell(uid).strip()

    with _WRITE_LOCK:
        target = None
        for a in load_accounts(path):
            if a.uid == uid:
                target = a
                break
        if target is None:
            raise CredError("找不到 UID 为 {} 的账号".format(uid))
        if target.enabled == enabled:
            return target               # 幂等: 没变化就不动文件，也不做备份

        wb = _load_for_write(path)
        ws = wb.worksheets[0]
        header_row, cols = _header_map(ws)
        status_col = _ensure_status_column(ws, header_row, cols)
        _clean_sheet(ws, header_row, cols)   # 同一个假行也会让状态修改写不进去
        ws.cell(row=target.row, column=status_col,
                value=STATUS_ENABLED if enabled else STATUS_DISABLED)

        _write_workbook(wb, path, backup_dir)

    for a in load_accounts(path):
        if a.uid == uid:
            return a
    raise CredError("写入成功但没读回该账号，请检查 {}".format(path))


if __name__ == "__main__":  # 自检: 只打印脱敏信息
    for i, acct in enumerate(load_accounts(), start=1):
        print("{}. {:24} UID={:12} AK={:24} {}".format(
            i, acct.email, acct.uid, acct.ak_masked,
            "启用" if acct.enabled else "停用"))
